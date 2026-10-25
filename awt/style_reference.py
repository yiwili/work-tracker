"""Historical records are references, never evidence for today's tasks."""
import base64
import csv
import io
import json
import re
import zipfile
from datetime import datetime
from xlrd.biffh import XLRDError
from openpyxl.utils.exceptions import InvalidFileException
from pathlib import Path
from sqlalchemy import select, func
from . import db, model_provider, prompts
from .collectors import stable
from .git_tools import git
from .privacy import redact

LIMIT = 20_000_000


def read_rows(payload):
    if payload.content_base64 is not None:
        try:
            data = base64.b64decode(payload.content_base64, validate=True)
        except ValueError:
            raise ValueError('上传文件编码无效') from None
        name = payload.filename
    else:
        path = Path(payload.path).expanduser()
        if not path.is_file() or path.stat().st_size > LIMIT:
            raise ValueError('请选择不超过20MB的历史文件（服务运行电脑上的路径）')
        data, name = path.read_bytes(), path.name
    if len(data) > LIMIT:
        raise ValueError('历史文件不能超过20MB')
    suffix = Path(name).suffix.lower()
    try:
        if suffix in ('.json', '.csv'):
            try:
                text = data.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = data.decode('gb18030')
            if suffix == '.json':
                rows = json.loads(text)
                if isinstance(rows, dict):
                    rows = next((rows[k] for k in ('items','data','rows') if isinstance(rows.get(k),list)), None)
            else:
                try:
                    dialect = csv.Sniffer().sniff(text[:10000], delimiters=',;\t')
                except csv.Error:
                    dialect = csv.excel
                rows = list(csv.DictReader(io.StringIO(text), dialect=dialect))
        elif suffix in ('.xlsx', '.xls'):
            if suffix == '.xlsx':
                from openpyxl import load_workbook
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    if sum(i.file_size for i in archive.infolist()) > 100_000_000:
                        raise ValueError('Excel解压体积过大')
                workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                try:
                    sheet=workbook.worksheets[0]
                    if sheet.max_row and sheet.max_row>50001:
                        raise ValueError('最多50000条记录')
                    values = list(sheet.iter_rows(values_only=True))
                finally:
                    workbook.close()
            else:
                import xlrd
                workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
                try:
                    sheet = workbook.sheet_by_index(0)
                    if sheet.nrows>50001:
                        raise ValueError('最多50000条记录')
                    values = [[xlrd.xldate_as_datetime(cell.value,workbook.datemode) if cell.ctype==xlrd.XL_CELL_DATE else cell.value for cell in sheet.row(i)] for i in range(sheet.nrows)]
                finally:
                    workbook.release_resources()
            rows = [dict(zip([str(h or '').strip() for h in values[0]], row)) for row in values[1:]] if values else []
        else:
            raise ValueError('支持 JSON、CSV、XLSX 和 XLS 文件')
        if not isinstance(rows,list) or len(rows)>50000:
            raise ValueError('文件应包含工作项列表，最多50000条')
        return rows
    except (ValueError, OSError, csv.Error, KeyError, IndexError, zipfile.BadZipFile, XLRDError, InvalidFileException) as exc:
        raise ValueError('无法读取历史文件，请检查格式、表头和文件大小：'+str(exc)[:150]) from None


def text_value(row, *keys):
    for key in keys:
        value = row.get(key)
        if value is not None and str(value).strip():
            if isinstance(value,dict):
                value = value.get('full_name') or value.get('name') or value.get('title') or value.get('path') or ''
            if isinstance(value,float) and value.is_integer():
                value=int(value)
            return str(value).strip()
    return ''


def work_date(value):
    for pattern in ('%Y-%m-%d','%Y/%m/%d','%m/%d/%Y','%Y年%m月%d日'):
        try:
            return datetime.strptime(value[:10] if pattern=='%Y-%m-%d' else value,pattern).date().isoformat()
        except ValueError:
            continue
    return ''


def ingest(conn, rows, since, source='file', label='', include_git=False):
    from .service import get_settings, now, save_settings, upsert
    settings = get_settings(conn)
    added = updated = skipped = 0
    for row in rows:
        if not isinstance(row,dict):
            skipped += 1
            continue
        title = text_value(row,'title','标题','工作项标题','工作项名称','任务标题','名称')
        if not title:
            skipped += 1
            continue
        record_date = work_date(text_value(row,'plan_started_at','work_date','date','日期','计划开始时间','开始日期','created_at','创建时间'))
        project = text_value(row,'project_path','project','所属项目','项目','program_name','repository')
        source_id = text_value(row,'id','number','编号','工作项编号')
        url = text_value(row,'html_url','url','链接')
        identity = stable('external',project,source_id) if source_id else stable('natural',project,record_date,title)
        value = {'id':identity,'title':redact(title,settings.redact_patterns)[:500],
                 'details':redact(text_value(row,'content','body','details','详细内容','描述'),settings.redact_patterns)[:20000],
                 'project':redact(project,settings.redact_patterns)[:500], 'work_date':record_date,
                 'source':source,'source_id':source_id[:200], 'url':url[:2000] if url.startswith('https://gitee.com/') else '', 'imported_at':now()}
        existed = conn.execute(select(db.history.c.id).where(db.history.c.id==identity)).first()
        upsert(conn,db.history,value,'id')
        updated += bool(existed)
        added += not bool(existed)
    selected = list(conn.execute(select(db.history).where(db.history.c.work_date>=str(since)).order_by(db.history.c.work_date,db.history.c.id)).mappings())
    examples = list(dict.fromkeys(r['title'] for r in selected))[-55:]
    git_titles=[]
    if include_git:
        roots=list(dict.fromkeys(settings.git_roots+list(conn.execute(select(db.repositories.c.root)).scalars())))
        for root in roots[:10]:
            for line in git(root,'log',f'--since={since}','--no-merges','--format=%s','-40').splitlines():
                title=redact(line.strip(),settings.redact_patterns)[:200]
                if title and title not in git_titles:
                    git_titles.append(title)
    settings.title_examples=list(dict.fromkeys(examples+git_titles[:25]))[:80]
    result={'file':label,'since':str(since),'records':len(rows),'selected_records':len(selected),
            'stored_records':conn.execute(select(func.count()).select_from(db.history)).scalar_one(),
            'added':added,'updated':updated,'skipped':skipped,'examples':len(settings.title_examples),'git_examples':len(git_titles[:25]),'imported_at':now()}
    settings.history_reference=result
    save_settings(conn,settings)
    return result


def import_reference(conn, payload):
    return ingest(conn,read_rows(payload),payload.since,label=payload.filename or payload.path,include_git=payload.include_git)


def generate_template(conn, force=False):
    from .service import get_settings,save_settings
    settings=get_settings(conn)
    if settings.style_template and not force:
        return {'generated':False,'message':'沿用已有风格模板','template':settings.style_template}
    if not settings.title_examples:
        raise ValueError('请先导入历史工作项')
    if not settings.llm_enabled:
        return {'generated':False,'message':'历史记录已保存；开启AI后可生成风格模板','template':settings.style_template}
    baseline=stable(settings.style_template,settings.title_examples,settings.agent_prompts)
    text=model_provider.complete(settings,prompts.prompt(settings,'style_template'),json.dumps({'titles':settings.title_examples},ensure_ascii=False),max_tokens=4000)
    try:
        value=json.loads(re.sub(r'^```(?:json)?\s*|\s*```$','',text.strip()))['template']
        if not isinstance(value,str) or not 1<=len(value.strip())<=12000:
            raise ValueError()
    except (ValueError,KeyError,TypeError):
        raise model_provider.ModelError('AI风格模板格式无效；历史数据已保留') from None
    conn.commit();conn.exec_driver_sql('BEGIN IMMEDIATE')
    current=get_settings(conn)
    if stable(current.style_template,current.title_examples,current.agent_prompts)!=baseline:
        raise ValueError('生成期间风格设置发生变化，已保留用户修改，请重试')
    current.style_template=redact(value.strip(),current.redact_patterns)
    save_settings(conn,current)
    return {'generated':True,'message':'已生成并启用历史命名风格模板','template':current.style_template}


def initial_template(conn):
    try:
        return generate_template(conn)
    except ValueError as exc:
        conn.rollback()
        return {'generated':False,'message':'历史已入库；'+str(exc),'template':''}


def listing(conn, offset=0, limit=20):
    return {'total':conn.execute(select(func.count()).select_from(db.history)).scalar_one(),
            'items':[dict(r) for r in conn.execute(select(db.history).order_by(db.history.c.work_date.desc(),db.history.c.id).offset(offset).limit(limit)).mappings()]}
